import sys,glob,docx
from pptx import Presentation
from pypdf import PdfReader
out=open('Website/scratch/all.txt','w')
for f in sorted(glob.glob('Presentations/*.pptx')):
    if '~$' in f: continue
    out.write(f'\n=====PPTX {f}\n')
    for i,s in enumerate(Presentation(f).slides):
        out.write(f'--slide {i+1}\n')
        for sh in s.shapes:
            if sh.has_text_frame: out.write(sh.text_frame.text+'\n')
for f in ['Team Report/L4 Team Report D3.docx','Team Report/Team Report - Design.docx']:
    out.write(f'\n=====DOCX {f}\n')
    d=docx.Document(f)
    for p in d.paragraphs:
        if p.text.strip(): out.write(p.text+'\n')
for f in glob.glob('Presentations/*.pdf'):
    if 'Todd' in f: continue
    out.write(f'\n=====PDF {f}\n')
    for p in PdfReader(f).pages: out.write((p.extract_text() or '')+'\n')
